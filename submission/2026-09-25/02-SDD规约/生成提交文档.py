#!/usr/bin/env python3
"""Build readable DOCX copies and a bounded Spec Markdown snapshot."""
import hashlib
import json
import re
import shutil
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
DELIVERY = BASE.parent
ROOT = DELIVERY.parent.parent


def font(run, *, code=False):
    run.font.name = 'Consolas' if code else 'Calibri'
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    if code:
        run.font.size = Pt(8)


def inline(paragraph, text):
    # Keep visible source paths and URLs; render emphasis without literal Markdown.
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        code = part.startswith('`') and part.endswith('`')
        run = paragraph.add_run(part[2:-2] if bold else part[1:-1] if code else part)
        run.bold = bold
        font(run, code=code)


def docx_from_markdown(source):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(10)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name in ('Title','Heading 1','Heading 2','Heading 3'):
        doc.styles[name].font.color.rgb = RGBColor.from_string('17365D')
        doc.styles[name]._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    inline(footer, 'Knowpipe · 课程提交材料 · ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    lines = source.read_text().splitlines(); i=0; in_code=False
    while i < len(lines):
        line=lines[i]
        if line.startswith('```'):
            in_code = not in_code; i+=1; continue
        if in_code:
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
            font(p.add_run(line),code=True); i+=1; continue
        if not line.strip():
            i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                values=[v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in values):rows.append(values)
                i+=1
            if rows:
                table=doc.add_table(rows=1,cols=max(map(len,rows)));table.style='Light Shading Accent 1'
                for r,values in enumerate(rows):
                    cells=table.rows[0].cells if r==0 else table.add_row().cells
                    for c,value in enumerate(values):
                        inline(cells[c].paragraphs[0],value)
                        if r==0:
                            for run in cells[c].paragraphs[0].runs:run.bold=True
                doc.add_paragraph()
            continue
        h=re.match(r'^(#{1,4})\s+(.+)',line)
        if h:
            p=doc.add_paragraph(style='Title' if len(h[1])==1 else 'Heading '+str(min(3,len(h[1])-1)))
            inline(p,h[2])
        elif line.startswith('- '):
            inline(doc.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\. ',line):
            inline(doc.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',line))
        else:
            inline(doc.add_paragraph(),line)
        i+=1
    target=source.with_suffix('.docx');doc.save(target)
    # Open the generated package again: catches corrupt ZIP/XML before delivery.
    reread=Document(target)
    return {'file':str(target.relative_to(DELIVERY)),'paragraphs':len(reread.paragraphs),'tables':len(reread.tables)}


def main():
    snapshot=BASE/'原始Spec快照';snapshot.mkdir(exist_ok=True)
    files=list((ROOT/'specs').rglob('*.md'))
    files += [ROOT/'.specify/memory/constitution.md',ROOT/'docs/sdd-workflow.md',ROOT/'project_goal_and_agent_prompt.md']
    manifest=[]
    for source in sorted(files):
        relative=source.relative_to(ROOT);target=snapshot/relative
        target.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(source,target)
        manifest.append({'source':str(relative),'sha256':hashlib.sha256(source.read_bytes()).hexdigest()})
    (snapshot/'清单.json').write_text(json.dumps({'captured_at_utc':datetime.now(timezone.utc).isoformat(),'files':manifest},ensure_ascii=False,indent=2)+'\n')
    wanted=[BASE/'需求规约.md',BASE/'设计规约.md',DELIVERY/'03-项目报告/项目报告.md',DELIVERY/'05-分组分工表/分组分工表.md']
    print(json.dumps({'documents':[docx_from_markdown(f) for f in wanted],'spec_files':len(manifest)},ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
