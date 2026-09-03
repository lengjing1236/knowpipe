#!/usr/bin/env python3
"""Build the course selection presentation from its slide-oriented Markdown source.

The converter intentionally supports the small semantic subset used by
`课程项目选题汇报.md`.  Layout directives select presentation-oriented visual
compositions instead of copying report paragraphs into text-heavy slides.
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(10, 18, 32)
PANEL = RGBColor(20, 33, 52)
PANEL_ALT = RGBColor(25, 42, 64)
WHITE = RGBColor(245, 248, 252)
MUTED = RGBColor(174, 190, 207)
TEAL = RGBColor(57, 213, 211)
BLUE = RGBColor(80, 145, 255)
GOLD = RGBColor(255, 191, 71)
GREEN = RGBColor(87, 204, 153)
RED = RGBColor(255, 107, 107)
PURPLE = RGBColor(170, 124, 255)
FONT = "Microsoft YaHei"


@dataclass
class Section:
    title: str
    body: list[str] = field(default_factory=list)


@dataclass
class SlideSpec:
    layout: str
    title: str
    sections: list[Section]
    quote: str = ""
    notes: str = ""
    plain: list[str] = field(default_factory=list)


def clean_md(text: str) -> str:
    text = text.strip()
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "").replace("__", "")
    return text


def parse_markdown(path: Path) -> tuple[dict[str, str], list[SlideSpec]]:
    source = path.read_text(encoding="utf-8").replace("\r\n", "\n")
    metadata: dict[str, str] = {}

    if source.startswith("---\n"):
        end = source.find("\n---\n", 4)
        if end != -1:
            for line in source[4:end].splitlines():
                if ":" in line:
                    key, value = line.split(":", 1)
                    metadata[key.strip()] = value.strip()
            source = source[end + 5 :]

    blocks = re.split(r"(?m)^---\s*$", source)
    specs: list[SlideSpec] = []

    for block in blocks:
        if not block.strip():
            continue
        layout_match = re.search(r"<!--\s*layout:\s*([\w-]+)\s*-->", block)
        layout = layout_match.group(1) if layout_match else "generic"
        notes_match = re.search(r"<!--\s*notes:\s*(.*?)-->", block, re.S)
        notes = clean_md(notes_match.group(1)) if notes_match else ""
        visible = re.sub(r"<!--.*?-->", "", block, flags=re.S)

        title_lines: list[str] = []
        sections: list[Section] = []
        plain: list[str] = []
        quote = ""
        current: Section | None = None

        for raw in visible.splitlines():
            line = raw.strip()
            if not line:
                continue
            if line.startswith("# "):
                title_lines.append(clean_md(line[2:]))
            elif line.startswith("## "):
                current = Section(clean_md(line[3:]))
                sections.append(current)
            elif line.startswith("> "):
                quote = clean_md(line[2:])
            elif line.startswith("- "):
                item = clean_md(line[2:])
                if current:
                    current.body.append(item)
                else:
                    plain.append(item)
            else:
                item = clean_md(line)
                if current:
                    current.body.append(item)
                else:
                    plain.append(item)

        specs.append(
            SlideSpec(
                layout=layout,
                title="\n".join(title_lines),
                sections=sections,
                quote=quote,
                notes=notes,
                plain=plain,
            )
        )

    return metadata, specs


def rect(slide, x, y, w, h, fill, radius: bool = True, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def textbox(
    slide,
    text,
    x,
    y,
    w,
    h,
    *,
    size=20,
    color=WHITE,
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0.05,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.runs[0]
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rich_lines(
    slide,
    lines,
    x,
    y,
    w,
    h,
    *,
    size=17,
    color=MUTED,
    bullet=True,
    spacing=7,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = 0
    frame.margin_bottom = 0
    for idx, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = ("• " if bullet else "") + line
        paragraph.space_after = Pt(spacing)
        paragraph.line_spacing = 1.05
        run = paragraph.runs[0]
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_background(slide):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    rect(slide, 0, 0, 0.11, 7.5, TEAL, radius=False)


def add_header(slide, title, number):
    textbox(slide, title, 0.7, 0.34, 11.8, 0.72, size=28, bold=True)
    rect(slide, 0.7, 1.12, 0.72, 0.05, TEAL, radius=False)
    textbox(
        slide,
        f"{number:02d}",
        12.25,
        0.37,
        0.45,
        0.32,
        size=11,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_quote(slide, quote):
    if not quote:
        return
    rect(slide, 0.7, 6.63, 11.95, 0.5, PANEL_ALT)
    textbox(
        slide,
        quote,
        0.92,
        6.71,
        11.5,
        0.3,
        size=13,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_card(slide, section, x, y, w, h, *, accent=TEAL, body_size=16):
    rect(slide, x, y, w, h, PANEL)
    rect(slide, x, y, 0.07, h, accent, radius=False)
    textbox(slide, section.title, x + 0.25, y + 0.2, w - 0.45, 0.45, size=19, bold=True)
    rich_lines(
        slide,
        section.body,
        x + 0.25,
        y + 0.82,
        w - 0.5,
        h - 1.0,
        size=body_size,
        bullet=len(section.body) > 1,
        spacing=6,
    )


def render_cover(slide, spec):
    add_background(slide)
    rect(slide, 9.65, 0, 3.68, 3.35, PANEL_ALT)
    rect(slide, 10.35, 0.05, 2.8, 2.8, BLUE)
    rect(slide, 10.95, 0.68, 1.6, 1.6, BG)
    rect(slide, 0.72, 1.25, 0.11, 3.8, TEAL, radius=False)
    textbox(slide, spec.title, 1.12, 1.18, 8.65, 2.45, size=30, bold=True, line_spacing=1.12)
    if spec.quote:
        textbox(slide, spec.quote, 1.15, 4.15, 8.5, 0.65, size=19, color=TEAL, bold=True)
    if spec.plain:
        textbox(slide, spec.plain[-1], 1.15, 5.55, 8.5, 0.42, size=14, color=MUTED)
    textbox(slide, "KNOWPIPE · COURSE PROJECT", 10.1, 6.75, 2.55, 0.3, size=10, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)


def render_three_cards(slide, spec, number, colors=(TEAL, BLUE, GOLD)):
    add_background(slide)
    add_header(slide, spec.title, number)
    count = max(1, len(spec.sections))
    gap = 0.25
    width = (11.95 - gap * (count - 1)) / count
    for idx, section in enumerate(spec.sections):
        add_card(slide, section, 0.7 + idx * (width + gap), 1.45, width, 4.82, accent=colors[idx % len(colors)], body_size=16)
    add_quote(slide, spec.quote)


def render_sources(slide, spec, number):
    add_background(slide)
    add_header(slide, spec.title, number)
    colors = (TEAL, BLUE)
    for idx, section in enumerate(spec.sections[:2]):
        add_card(slide, section, 0.7 + idx * 6.1, 1.45, 5.85, 4.82, accent=colors[idx], body_size=15)
    add_quote(slide, spec.quote)


def render_architecture(slide, spec, number):
    add_background(slide)
    add_header(slide, spec.title, number)
    colors = (PURPLE, BLUE, TEAL, GREEN, GOLD)
    x = 0.7
    y = 2.05
    w = 2.18
    gap = 0.25
    for idx, section in enumerate(spec.sections):
        add_card(slide, section, x + idx * (w + gap), y, w, 3.25, accent=colors[idx], body_size=13)
        if idx < len(spec.sections) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(x + idx * (w + gap) + w - 0.02),
                Inches(3.29),
                Inches(gap + 0.07),
                Inches(0.38),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = MUTED
            arrow.line.color.rgb = MUTED
    textbox(slide, "数据输入", 0.78, 1.55, 2.0, 0.3, size=12, color=MUTED)
    textbox(slide, "用户价值", 10.95, 5.52, 1.55, 0.3, size=12, color=MUTED, align=PP_ALIGN.RIGHT)


def render_mining(slide, spec, number):
    add_background(slide)
    add_header(slide, spec.title, number)
    colors = (PURPLE, BLUE, TEAL, GREEN)
    for idx, section in enumerate(spec.sections):
        x = 0.7 + idx * 3.03
        rect(slide, x, 1.62, 2.78, 4.55, PANEL)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(1.88), Inches(0.62), Inches(0.62))
        circle.fill.solid()
        circle.fill.fore_color.rgb = colors[idx]
        circle.line.color.rgb = colors[idx]
        textbox(slide, str(idx + 1), x + 0.2, 1.95, 0.62, 0.32, size=16, color=BG, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, section.title, x + 0.2, 2.75, 2.38, 0.5, size=20, bold=True)
        rich_lines(slide, section.body, x + 0.2, 3.45, 2.38, 1.9, size=15, bullet=False)
    add_quote(slide, spec.quote)


def render_personalization(slide, spec, number):
    add_background(slide)
    add_header(slide, spec.title, number)
    if not spec.sections:
        return
    knowledge = spec.sections[0]
    rect(slide, 0.7, 1.42, 11.95, 1.0, PANEL_ALT)
    textbox(slide, knowledge.title, 0.98, 1.65, 2.2, 0.35, size=18, bold=True, color=TEAL)
    textbox(slide, " · ".join(knowledge.body), 3.0, 1.63, 9.0, 0.38, size=16, color=WHITE)
    colors = (MUTED, GOLD, GREEN)
    for idx, section in enumerate(spec.sections[1:4]):
        add_card(slide, section, 0.7 + idx * 4.07, 2.75, 3.8, 3.5, accent=colors[idx], body_size=15)
    add_quote(slide, spec.quote)


def render_two_columns(slide, spec, number, colors=(TEAL, BLUE)):
    add_background(slide)
    add_header(slide, spec.title, number)
    for idx, section in enumerate(spec.sections[:2]):
        add_card(slide, section, 0.7 + idx * 6.1, 1.45, 5.85, 4.82, accent=colors[idx], body_size=15)
    add_quote(slide, spec.quote)


def render_bands(slide, spec, number):
    add_background(slide)
    add_header(slide, spec.title, number)
    colors = (GREEN, BLUE, PURPLE)
    for idx, section in enumerate(spec.sections):
        y = 1.42 + idx * 1.63
        rect(slide, 0.7, y, 11.95, 1.32, PANEL)
        rect(slide, 0.7, y, 1.75, 1.32, colors[idx], radius=False)
        textbox(slide, section.title, 0.85, y + 0.39, 1.42, 0.42, size=19, color=BG, bold=True, align=PP_ALIGN.CENTER)
        textbox(slide, " · ".join(section.body), 2.75, y + 0.33, 9.55, 0.58, size=17, color=WHITE, valign=MSO_ANCHOR.MIDDLE)
    add_quote(slide, spec.quote)


def render_timeline(slide, spec, number):
    add_background(slide)
    add_header(slide, spec.title, number)
    colors = (BLUE, TEAL, GOLD)
    for idx, section in enumerate(spec.sections):
        x = 0.7 + idx * 4.07
        rect(slide, x, 1.65, 3.8, 4.55, PANEL)
        rect(slide, x, 1.65, 3.8, 0.14, colors[idx], radius=False)
        textbox(slide, f"0{idx + 1}", x + 0.25, 2.05, 0.7, 0.42, size=20, color=colors[idx], bold=True)
        textbox(slide, section.title, x + 0.25, 2.75, 3.25, 0.48, size=21, bold=True)
        rich_lines(slide, section.body, x + 0.25, 3.48, 3.25, 1.8, size=15, bullet=False)
    add_quote(slide, spec.quote)


def render_grid(slide, spec, number, colors=(TEAL, BLUE, GOLD, PURPLE)):
    add_background(slide)
    add_header(slide, spec.title, number)
    for idx, section in enumerate(spec.sections[:4]):
        col = idx % 2
        row = idx // 2
        add_card(
            slide,
            section,
            0.7 + col * 6.1,
            1.43 + row * 2.42,
            5.85,
            2.15,
            accent=colors[idx],
            body_size=14,
        )
    add_quote(slide, spec.quote)


def render_close(slide, spec):
    add_background(slide)
    textbox(slide, spec.title, 0.8, 0.8, 11.7, 0.75, size=34, bold=True, align=PP_ALIGN.CENTER)
    main = spec.sections[0] if spec.sections else Section("", spec.plain)
    textbox(slide, main.title, 1.2, 1.9, 10.9, 0.62, size=23, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    colors = (PURPLE, BLUE, TEAL, GREEN, GOLD)
    items = main.body or spec.plain
    for idx, item in enumerate(items[:5]):
        x = 0.76 + idx * 2.5
        rect(slide, x, 3.15, 2.25, 1.35, PANEL)
        rect(slide, x, 3.15, 2.25, 0.1, colors[idx], radius=False)
        textbox(slide, item, x + 0.13, 3.52, 1.99, 0.55, size=16, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_quote(slide, spec.quote)
    textbox(slide, "选题确认 · 等待人工评审", 4.2, 5.55, 4.95, 0.42, size=15, color=MUTED, align=PP_ALIGN.CENTER)


def render_generic(slide, spec, number):
    add_background(slide)
    add_header(slide, spec.title, number)
    body = spec.plain[:]
    for section in spec.sections:
        body.append(section.title)
        body.extend(section.body)
    rich_lines(slide, body, 0.9, 1.55, 11.4, 4.8, size=18)
    add_quote(slide, spec.quote)


def add_notes(slide, notes):
    if not notes:
        return
    try:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = notes
    except Exception:
        # Notes support varies across python-pptx versions; slide generation must
        # remain deterministic even if a version cannot write speaker notes.
        pass


def build(input_path: Path, output_path: Path):
    metadata, specs = parse_markdown(input_path)
    if not specs:
        raise ValueError(f"No slides found in {input_path}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    prs.core_properties.title = metadata.get("title", specs[0].title.replace("\n", " "))
    prs.core_properties.subject = metadata.get("description", "课程项目选题确认汇报")
    prs.core_properties.author = "陶文杰, 陈铭宇"
    prs.core_properties.comments = "Generated from 课程项目选题汇报.md"

    for number, spec in enumerate(specs, start=1):
        slide = prs.slides.add_slide(blank)
        if spec.layout == "cover":
            render_cover(slide, spec)
        elif spec.layout in {"problem", "statement"}:
            render_three_cards(slide, spec, number)
        elif spec.layout == "sources":
            render_sources(slide, spec, number)
        elif spec.layout == "architecture":
            render_architecture(slide, spec, number)
        elif spec.layout == "mining":
            render_mining(slide, spec, number)
        elif spec.layout == "personalization":
            render_personalization(slide, spec, number)
        elif spec.layout in {"runtime", "owners"}:
            colors = (BLUE, GOLD) if spec.layout == "runtime" else (TEAL, PURPLE)
            render_two_columns(slide, spec, number, colors)
        elif spec.layout == "scope":
            render_bands(slide, spec, number)
        elif spec.layout == "timeline":
            render_timeline(slide, spec, number)
        elif spec.layout in {"gates", "evidence"}:
            render_grid(slide, spec, number)
        elif spec.layout == "close":
            render_close(slide, spec)
        else:
            render_generic(slide, spec, number)
        add_notes(slide, spec.notes)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return len(specs)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="课程项目选题汇报.md", type=Path)
    parser.add_argument("--output", default="课程项目选题汇报.pptx", type=Path)
    args = parser.parse_args()
    slide_count = build(args.input, args.output)
    print(f"Generated {args.output} with {slide_count} slides")


if __name__ == "__main__":
    main()
