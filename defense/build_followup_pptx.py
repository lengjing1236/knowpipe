"""Generate an evidence-backed progress deck; do not present pending gates as completed."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parent.parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
INK, GREEN = RGBColor(32, 52, 46), RGBColor(36, 91, 69)


def text(slide, content, x, y, w, h, size=23, color=INK):
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    frame.word_wrap = True
    for index, line in enumerate(content.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = 'Microsoft YaHei'
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(16)


def slide(title, lines, note=None, picture=None):
    page = prs.slides.add_slide(prs.slide_layouts[6])
    page.background.fill.solid()
    page.background.fill.fore_color.rgb = RGBColor(245, 246, 240)
    text(page, 'KNOWPIPE  /  2026.09.17', .65, .28, 12, .35, 11, GREEN)
    text(page, title, .65, .85, 12, .9, 30, GREEN)
    if picture:
        page.shapes.add_picture(str(ROOT / picture), Inches(.7), Inches(1.9), height=Inches(4.9))
        text(page, '\n'.join(lines), 6.6, 2, 6, 4.7, 21)
    else:
        text(page, '\n'.join(lines), .8, 2, 11.7, 4.6)
    text(page, f'答辩演示 · 已验证与待验收分开记录                           {len(prs.slides)}', .7, 7.02, 12, .3, 10)
    if note:
        page.notes_slide.notes_text_frame.text = note


slide('从订阅更新，到可追溯的知识产物', [
    '技术文献 + 播客文字稿：多源知识发现与学习辅助',
    'Spark 负责核心文本计算，MongoDB 保存数据与处理证据',
    '本次交付：可运行 Web、真实编程播客分析、通知与评价工具',
    '尚待完成：永久云部署、万条最终批次、真实人工评价'])
slide('四层架构与进程边界', [
    '交互层：响应式网页、文字稿阅读、SSE 与轮询通知',
    '业务层：账户、画像、订阅、任务状态与结果查询',
    '数据层：MongoDB 文档、唯一索引、批次、持久通知',
    '计算层：独立 worker + Spark TF-IDF / KMeans / 相似度',
    '演示：本机 Gunicorn + 可选免费 HTTPS 隧道；计算独立运行'])
slide('用户可操作的知识工作台', [
    '登录后自动恢复身份', '订阅动态与缺失文字稿状态', '手机页面无横向溢出',
    '数字来自真实数据库', '截图为 JS Party 真实节目，文献数为 0'], picture='evidence/004-podcast-insights/local-real-podcast-desktop.png')
slide('真实业务验收：JS Party 编程播客', [
    '公开 RSS：changelog.com/jsparty/feed（约 2.27 MB）',
    'React: then & now：官方文字稿 69,702 字符，88 个片段',
    'WYSIWYG：官方文字稿 75,171 字符，94 个片段',
    '真实 Spark local[2] → MongoDB → Web；无文字稿合成或 fetch 注入',
    '历史节目重复检查不重算、不重复通知；未测新发布延迟'])
slide('文字稿加工：每一步都能解释', [
    '公开官方文字稿，或用户补充的文本',
    '中英文分词 → Spark TF-IDF → 主题分组 → 相似片段',
    '关键词权重、原文片段、批次和 application id 可查看',
    '单段/重复内容显式词频降级，不伪造多主题结果',
    '没有文字稿时保持等待；当前不自动转写音频'])
slide('分钟级增量处理与通知', [
    '默认每 5 分钟检查 RSS；SSE 推送 + 每 5 秒查询持久通知',
    '首次回填 3 集，后续检查更新；不继续导入更早的历史全集',
    '唯一键去重、最多三次自动重试、租约过期恢复',
    '每用户每节目一条通知，订阅与通知访问相互隔离',
    '这是定期增量批处理，不是直播流处理'])
slide('Spark 改进与算法边界', [
    '文献相似度从 driver 循环迁移到 Spark task',
    '仍采用有界候选组：不是全库精确近邻',
    '播客相似度通过 Spark SQL 稀疏词项连接与聚合计算',
    '批次记录 Spark master / application id',
    '本轮验证 local[2]；公网部署不代表多机集群'])
slide('验证证据与规约驱动过程', [
    'codegraph 上下文分析 → spec / plan / tasks / 契约',
    '先行失败测试 → 实现 → 自动验证 → 验收记录',
    '99 项完整测试通过，Spark 测试使用真实 Java / PySpark',
    '真实 Mongo + 双进程 Gunicorn + 浏览器流程通过',
    '新增真实播客验收：2 集完成、182 段、2 条通知，1 集等待'])
slide('推荐效果：需要真实人工标签', [
    '为不同用户画像构造个性化与对照基线排名',
    '同一文档统一 ID，候选合并后进行盲标 0–3 相关性',
    '工具计算 Precision@K、NDCG@K 和差值',
    '未标注拒绝计算；短列表不提高分母优势',
    '工具已验证，真实评价尚未完成，不宣称质量提升'])
slide('下一步与最终验收门槛', [
    '明日答辩：本机入口 + 免费临时 HTTPS；永久云部署后续补齐',
    '最终数据：至少 10,000 条有效 Stack Exchange + arXiv',
    '播客：历史真实节目已验收；新节目到达与自动转写后续扩展',
    '真实人工评价：至少两种画像并保留标注依据',
    '后续优先：跨节目趋势、播客与技术文献关联'])
prs.save(ROOT / 'defense/答辩演示-20260917.pptx')
print(f'Saved {len(prs.slides)} slides; pending gates explicitly labeled.')
