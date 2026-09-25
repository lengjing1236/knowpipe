"""Build the six-deliverable index and a credential-free course submission zip."""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import subprocess
import zipfile
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / 'submission/2026-09-25'


def git(*args):
    return subprocess.check_output(['git', *args], cwd=ROOT, text=True).strip()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--package', action='store_true')
    args = parser.parse_args()
    source = DEST / '01-源代码'
    source.mkdir(parents=True, exist_ok=True)
    branch, revision = git('branch', '--show-current'), git('rev-parse', 'HEAD')
    (source / '源码与运行说明.md').write_text(f'''# 源代码与本机演示

Git仓库：[Knowpipe](https://github.com/lengjing1236/knowpipe)。提交分支：`{branch}`。
打包基线提交：`{revision}`。远程推送状态见总清单；保留真实历史，不改写提交时间。

## 这台电脑直接使用

```bash
python3 scripts/demo_release.py status
# 服务未运行时：
python3 scripts/demo_release.py start
```

打开 http://127.0.0.1:8019/learning 。用户名和本机随机密码保存在`state/demo-release/account.json`，该私有文件不进Git或提交包。

```bash
cat state/demo-release/account.json
```

演示目标：学习 Python 的异常捕获、异常传播和日志记录。
推荐由真实后台计算，首次准备可能需要数分钟；录屏等待如有剪辑会标注。可直接查看当前已计算结果，也可保存新目标、标记/撤销已读并观察重算。

## 从源码恢复

环境为Python 3.10、Java 17、MongoDB；安装`requirements.txt`，核心程序在`knowpipe/`。本次使用MongoDB 127.0.0.1:27018及隔离库knowpipe_demo012。
首次执行`python3 scripts/demo_release.py prepare`会读取本机已有万条语料中的全部12篇中文全文。提交目录的`演示中文全文.jsonl`提供同一中文子集；换机可用`--corpus`指定它，实际参数见脚本`--help`。
本机语义模型是排序/相似度组件，准备方法见`specs/011-mvp-recommendation-validation/quickstart.md`与`prepare_semantic011.py --help`；翻译只用远程免费服务。模型权重、MongoDB数据目录、音频及万条原文不包含在源码快照中。

## 免费远程翻译

固定使用GLM-4.7-Flash。配置`KNOWPIPE_BIGMODEL_API_KEY`，或将密钥保存在私有`state/demo012/bigmodel-api-key`（权限0600），然后stop/start重启演示。
密钥不得放入提交包。缺少密钥时原生中文仍可阅读，英文翻译显示未配置；不能回落本地或付费模型。真实调用情况见`evidence/012-demo-delivery/remote-translation.json`（如存在）。

## 操作顺序

1. 登录，保存中文学习目标。
2. 等待推荐，展开“为什么推荐这篇”查看真实正文依据。
3. 阅读中文全文，主动“标记已读”。
4. 查看阅读记录和新版推荐；已读不等于已经掌握。
5. 用资料库筛选来源；RSS入口及历史音频处理证据另行展示。

边界：本机单机Spark；演示12篇中文子集与10215篇背景实验分开。011总体推荐/补充/中文质量未通过，视频不能代替该结论。
''', encoding='utf-8')
    (source / '真实Git提交记录.txt').write_text(git('log', '-100', '--date=iso-strict', '--pretty=format:%h | %ad | %s')+'\n',encoding='utf-8')
    corpus=ROOT/'state/feature009/corpus-expanded/documents.jsonl'
    if corpus.exists():
        records=[]
        for line in corpus.open(encoding='utf-8'):
            item=json.loads(line)
            if str(item.get('language','')).lower().startswith('zh'):
                records.append(item)
        (source/'演示中文全文.jsonl').write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in records),encoding='utf-8')
        (source/'资料来源说明.md').write_text('# 演示资料出处\n\n本文件收录既有采集的全部中文全文，保留来源URL、版本与许可字段。资料作者及许可归原发布方；这些资料不是本项目原创教材。\n\n'+'\n'.join(f'- [{r["title"]}]({r["source_url"]})；许可：{r.get("license", "见原始页面")}' for r in records)+'\n',encoding='utf-8')
    if not args.package:
        print('源码说明及可移植中文演示子集已生成；最终包等待真实视频/PPT。')
        return
    report=json.loads((ROOT/'evidence/012-demo-delivery/browser.json').read_text())
    assert report['status']=='passed' and float(report['video']['duration'])<=300
    required={
        '01源代码':source/'源码与运行说明.md',
        '02SDD规约':DEST/'02-SDD规约/需求规约.docx',
        '02设计规约':DEST/'02-SDD规约/设计规约.docx',
        '02提示词':DEST/'02-SDD规约/提示词归档/真实用户提示词.md',
        '03项目报告':DEST/'03-项目报告/项目报告.docx',
        '04视频':ROOT/report['video']['path'],
        '05分工表':DEST/'05-分组分工表/分组分工表.docx',
        '06PPT':DEST/'06-汇报PPT/Knowpipe项目汇报.pptx'}
    for label,path in required.items():
        if not path.is_file() or not path.stat().st_size:
            raise AssertionError(f'缺少交付物：{label}')
        if path.suffix=='.docx': Document(path)
        if path.suffix=='.pptx':
            presentation=Presentation(path)
            if len(presentation.slides)<6: raise AssertionError('PPT页数异常')
    evidence=DEST/'核验依据'
    evidence.mkdir(exist_ok=True)
    for name in ['browser.json','acceptance.md','remote-translation.json','remote-translation-diagnostic.json']:
        p=ROOT/'evidence/012-demo-delivery'/name
        if p.exists(): (evidence/name).write_bytes(p.read_bytes())
    old=ROOT/'evidence/011-mvp-recommendation-validation/acceptance.md'
    (evidence/'011算法质量验收.md').write_bytes(old.read_bytes())
    archive=source/'源码快照.zip'
    with archive.open('wb') as stream:
        subprocess.run(['git','archive','--format=zip','HEAD','knowpipe','scripts','tests','requirements.txt','requirements','README.md','specs','docs'],cwd=ROOT,stdout=stream,check=True)
    remote_file=ROOT/'evidence/012-demo-delivery/git-remote.json'
    remote=json.loads(remote_file.read_text()) if remote_file.exists() else {'status':'待核验'}
    manifest={'created_at':datetime.now(timezone.utc).isoformat(),'source_revision':revision,'branch':branch,
              'remote':remote,'browser_status':report['status'],'video_seconds':float(report['video']['duration']),
              'files':[]}
    for path in sorted(DEST.rglob('*')):
        if path.is_file() and path.name not in ('文件清单.json','提交说明.md'):
            data=path.read_bytes()
            manifest['files'].append({'path':str(path.relative_to(DEST)),'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()})
    (DEST/'文件清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    (DEST/'提交说明.md').write_text(f'''# Knowpipe 六项课程交付

打开本目录对应编号的文件即可提交；不要提交state目录或任何API密钥。

| 序号 | 交付物 | 文件 |
| --- | --- | --- |
| 1 | 源代码和真实Git历史 | 01-源代码/源码与运行说明.md、源码快照.zip；[Git仓库](https://github.com/lengjing1236/knowpipe/tree/{branch}) |
| 2 | SDD需求/设计及提示词 | 02-SDD规约/，Word与Markdown；提示词归档注明可访问记录范围 |
| 3 | 项目报告 | 03-项目报告/项目报告.docx |
| 4 | 必交备份演示视频 | 04-备份演示视频/Knowpipe演示.mp4（{float(report['video']['duration']):.1f}秒，中文字幕） |
| 5 | 最终分工表 | 05-分组分工表/分组分工表.docx；本人真实信息以确认栏为准 |
| 6 | 汇报PPT | 06-汇报PPT/Knowpipe项目汇报.pptx |

网页入口：http://127.0.0.1:8019/learning 。登录信息仅在本机state/demo-release/account.json。

## 实际范围

012录制完成真实目标、推荐、全文阅读、主动已读与重算。演示是12篇原生中文子集，完整背景10215篇的效果问题仍见核验依据/011算法质量验收.md。远程翻译实测状态以核验依据中的相应文件为准。RSS音频到个人通知闭环尚未完整通过，不在视频中伪造通知。

## 提交前仅需核对

分工已按用户确认填为两人各50%，与初版无调整；如学校另行要求，补充学号班级并核对指定提交渠道。本工具没有向教师或群聊发送材料。

源码基线：{revision}。远程状态：{remote.get('status','待核验')}。文件完整性见文件清单.json。
''',encoding='utf-8')
    target=ROOT/'submission/Knowpipe-六项交付-2026-09-25.zip'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for path in sorted(DEST.rglob('*')):
            if path.is_file(): z.write(path,str(path.relative_to(DEST.parent)))
    with zipfile.ZipFile(target) as z:
        assert z.testzip() is None
    print(json.dumps({'package':str(target),'bytes':target.stat().st_size,'files':len(manifest['files'])},ensure_ascii=False))


if __name__=='__main__': main()
