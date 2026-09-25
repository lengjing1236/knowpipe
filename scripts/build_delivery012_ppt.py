#!/usr/bin/env python3
"""Build an editable Chinese presentation plus matching PIL layout previews.

Previews reproduce our drawing primitives, not PowerPoint's rendering engine.
Final mode requires the four real browser screenshots; --draft names a draft.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'submission/2026-09-25/06-汇报PPT'
EVIDENCE = ROOT/'evidence/012-demo-delivery'
FONT = '/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf'
LATIN = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
W,H,SCALE = 1600,900,120
BG='#F5F3E9'; INK='#183F32'; MUTED='#607268'; GREEN='#237652'; PALE='#DEEADD'; WHITE='#FFFFFF'; WARN='#A55E30'; RED='#AC514A'; LINE='#C9D5C8'


def color(value):
    return RGBColor.from_string(value.lstrip('#'))


class MixedFont:
    def __init__(self,size):
        self.cjk=ImageFont.truetype(FONT,size);self.latin=ImageFont.truetype(LATIN,size)

    def choose(self,char):
        n=ord(char)
        return self.cjk if (0x2E80<=n<=0x9FFF or 0xF900<=n<=0xFAFF or 0xFF00<=n<=0xFFEF) else self.latin

    def getlength(self,text):
        return sum(self.choose(char).getlength(char) for char in text)

    def draw(self,canvas,xy,text,fill):
        x,y=xy
        for char in text:
            font=self.choose(char);canvas.text((x,y),char,font=font,fill=fill,anchor='lt')
            x+=font.getlength(char)


def wrap(text, width, font):
    output=[]
    for original in text.split('\n'):
        line=''
        for char in original:
            if line and font.getlength(line+char)>width:
                output.append(line);line=char
            else:line+=char
        output.append(line)
    return output


class Deck:
    def __init__(self,draft=False):
        self.prs=Presentation();self.prs.slide_width=Inches(W/SCALE);self.prs.slide_height=Inches(H/SCALE)
        self.images=[];self.draft=draft;self.screenshot_manifest=[]

    def page(self,title,kicker='KNOWPIPE  /  课程项目汇报',note=''):
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=color(BG)
        self.image=Image.new('RGB',(W,H),BG);self.draw=ImageDraw.Draw(self.image)
        self.images.append(self.image)
        self.rect(0,0,18,H,GREEN)
        self.text(66,36,1440,36,kicker,12,MUTED)
        self.text(66,94,1465,106,title,30,INK,True)
        self.rect(66,190,1460,2,LINE)
        self.text(66,842,1400,30,'单机 Spark · 实际网页与数据 · 演示运行和推荐质量分别验收',11,MUTED)
        self.text(1450,842,90,30,f'{len(self.images):02d}',12,GREEN,True)
        self.slide.notes_slide.notes_text_frame.text=note
        return self

    def rect(self,x,y,w,h,fill=WHITE,line=None,round=False):
        shape=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,
                                         Inches(x/SCALE),Inches(y/SCALE),Inches(w/SCALE),Inches(h/SCALE))
        shape.fill.solid();shape.fill.fore_color.rgb=color(fill)
        if line:shape.line.color.rgb=color(line)
        else:shape.line.fill.background()
        if round:self.draw.rounded_rectangle((x,y,x+w,y+h),radius=18,fill=fill,outline=line)
        else:self.draw.rectangle((x,y,x+w,y+h),fill=fill,outline=line)
        return shape

    def text(self,x,y,w,h,value,size=20,fill=INK,bold=False):
        pil_font=MixedFont(round(size*SCALE/72))
        lines=wrap(value,w-4,pil_font)
        line_height=round(size*SCALE/72*1.33)
        if len(lines)*line_height>h+8:
            raise ValueError(f'text_overflow slide={len(self.images)} box={value[:50]!r} lines={len(lines)} h={h}')
        box=self.slide.shapes.add_textbox(Inches(x/SCALE),Inches(y/SCALE),Inches(w/SCALE),Inches(h/SCALE))
        tf=box.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.word_wrap=False
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
            p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color(fill)
            p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=Pt(size*1.33)
            pil_font.draw(self.draw,(x,y+i*line_height),line,fill)
        return box

    def card(self,x,y,w,h,title,body,accent=GREEN):
        self.rect(x,y,w,h,WHITE,round=True);self.rect(x,y,7,h,accent)
        self.text(x+26,y+22,w-48,90,title,20,accent,True)
        self.text(x+26,y+103,w-48,h-118,body,18,INK)

    def screenshot(self,name,x,y,w,h):
        path=EVIDENCE/name
        if not path.exists():
            if not self.draft:raise FileNotFoundError(path)
            self.rect(x,y,w,h,PALE,LINE,True)
            self.text(x+25,y+30,w-50,h-45,'草稿：等待实际页面截图\n'+name,20,MUTED)
            self.screenshot_manifest.append({'file':name,'exists':False})
            return
        img=Image.open(path).convert('RGB')
        fitted=ImageOps.contain(img,(int(w),int(h)))
        left=x+(w-fitted.width)/2;top=y+(h-fitted.height)/2
        self.rect(x,y,w,h,WHITE,LINE)
        self.slide.shapes.add_picture(str(path),Inches(left/SCALE),Inches(top/SCALE),width=Inches(fitted.width/SCALE),height=Inches(fitted.height/SCALE))
        self.image.paste(fitted,(int(left),int(top)))
        self.screenshot_manifest.append({'file':name,'exists':True,'pixels':list(img.size)})

    def save(self):
        OUT.mkdir(parents=True,exist_ok=True)
        name='Knowpipe项目汇报'+('-草稿' if self.draft else '')
        ppt=OUT/(name+'.pptx');self.prs.save(ppt)
        preview=OUT/('预览-草稿' if self.draft else '预览');preview.mkdir(exist_ok=True)
        for i,img in enumerate(self.images,1):img.save(preview/f'{i:02d}.png')
        thumb=Image.new('RGB',(960,((len(self.images)+2)//3)*180),BG)
        for i,img in enumerate(self.images):thumb.paste(img.resize((320,180)),((i%3)*320,(i//3)*180))
        thumb.save(preview/'总览.png')
        self.images[0].save(OUT/(name+'-布局预览.pdf'),save_all=True,append_images=self.images[1:],resolution=144)
        report={'pptx':str(ppt.relative_to(ROOT)),'slides':len(self.images),'draft':self.draft,
                'preview_note':'PIL按同一布局绘制，非PowerPoint/LibreOffice逐像素渲染；PPT文字与图形仍可编辑。',
                'screenshots':self.screenshot_manifest,'quality_verdict':'011 MVP未通过；012工程演示单独记录'}
        (OUT/(name+'-校验.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
        reopened=Presentation(ppt)
        assert len(reopened.slides)==len(self.images)
        assert all(s.notes_slide.notes_text_frame.text.strip() for s in reopened.slides)
        print(json.dumps(report,ensure_ascii=False))


def build(draft=False):
    d=Deck(draft)
    d.page('Knowpipe：从技术资料中找到下一步阅读',note='本项目面向有基础的计算机学生，按明确目标组织多源全文。今天展示真实可操作流程，同时报告推荐质量仍存在的缺口。')
    d.text(72,228,1410,120,'基于 Spark 与 MongoDB 的多源技术资料分析与个性化推荐 Web 系统',26,INK,True)
    d.card(72,374,456,242,'12 篇中文全文','本次实际网页演示\nPython 7篇 + Django 5篇')
    d.card(570,374,456,242,'10,215 篇背景','独立的规模与质量验证\n真实语料，不与演示混淆')
    d.card(1068,374,456,242,'一个核心问题','已读之后，下一篇资料\n能否直接帮助当前目标？')
    d.text(75,683,1400,80,'陶文杰 / 陈铭宇   ·   2026-09-25   ·   当前部署：单台电脑',20,MUTED)
    d.text(75,760,1400,50,'工程流程可以演示；011核心推荐质量尚未通过。',20,WARN,True)

    d.page('用户流程：目标 → 材料 → 阅读 → 再推荐',note='不是让系统判断学生已经掌握什么。用户主动保存目标和标记已读，系统据此重新处理；打开页面或读通知都不自动当作学会。')
    labels=[('01','输入中文目标','围绕具体技术问题'),('02','查找多源全文','先判断正文是否相关'),('03','给出原文依据','展示重合/可能补充'),('04','阅读中文全文','保留来源和处理状态'),('05','主动标记已读','记录本次正文版本'),('06','再次计算推荐','历史确实进入处理')]
    for i,(n,title,body) in enumerate(labels):
        x=72+(i%3)*496;y=236+(i//3)*266
        d.card(x,y,452,226,n+'  '+title,body)
    d.text(80,786,1400,42,'已读 ≠ 掌握；不同来源 ≠ 必然有新知识。',18,WARN,True)

    d.page('四层架构：Spark真实参与，长任务交给worker',note='四层按职责划分，不声称四套微服务或多主机部署。Spark做索引、召回、聚合和选择，Mongo保存事实，模型提供有界语义特征。增加Spark节点只能改变容量和吞吐，不能自动修复错误判断。')
    layers=[('Web交互层','中文目标 · 推荐列表 · 全文阅读 · 主动已读'),('业务层','Flask鉴权 · 任务队列 · worker租约 · 翻译/通知门控'),('数据管理与挖掘层','MongoDB原文/版本/历史 · 有界语义特征 · 选择规则'),('Spark处理层','全文索引 · 全库召回 · 特征聚合 · 覆盖/补充/去重')]
    for i,(title,body) in enumerate(layers):
        y=230+i*138;d.rect(74,y,992,114,PALE if i%2==0 else WHITE,round=True)
        d.text(98,y+17,920,41,title,20,GREEN,True);d.text(98,y+64,920,42,body,17,INK)
    d.card(1104,232,414,258,'本次：单机 Spark','local模式真实执行\n不是跨主机集群\nWeb不在HTTP内重算')
    d.card(1104,522,414,258,'先正确，再扩容','扩容可改善容量/吞吐\n不会自动提高相关性\n也不能修复翻译含义',WARN)

    d.page('多源数据：规模证据和演示数据分别说明',note='完整背景10167篇来自Stack Exchange，48篇来自四类官方文档。规模符合课程实践方向，但来源分布不均。当前12篇原生中文全文只是演示集合，并非万条都已译成中文。arXiv是历史采集能力，不在这个10215背景清单中。')
    counts=[('Stack Exchange 技术问答','10,167'),('PostgreSQL 官方文档','24'),('Docker 官方文档','12'),('Python 官方文档','7'),('Django 官方文档','5')]
    for i,(name,value) in enumerate(counts):
        y=235+i*98;d.rect(75,y,900,82,WHITE,round=True);d.text(102,y+21,670,45,name,21);d.text(778,y+18,170,48,value,23,GREEN,True)
    d.card(1018,235,500,240,'背景合计 10,215','用于011冻结质量评价\n来源与正文版本可追溯')
    d.card(1018,509,500,260,'本次演示 12 篇','Python 7 + Django 5\n原生中文官方全文\n不冒充全库中文翻译')
    d.text(80,777,1450,43,'来源：009 scale.json / corpus manifest。镜像平台不重复计作独立内容来源。',14,MUTED)

    d.page('核心算法：相关性先于“与已读不同”',note='先用Spark词汇召回，再用有界语义特征判断相关性及上下文。覆盖、补充和列表重复实际进入Spark选择。模型分数不是事实标签；无法比较时显示不确定，不能把中性预测等同新知识。')
    steps=[('1  全文召回','Spark词汇索引\n技术对象与正文过滤'),('2  有界语义比较','原中文目标相关性\n候选与已读上下文'),('3  聚合并选择','相关度 + 可能补充\n− 已读覆盖 − 列表重复'),('4  返回证据','双方原文、版本、位置\n不确定范围明确可见')]
    for i,(title,body) in enumerate(steps):d.card(75+i*370,246,338,298,title,body)
    d.rect(75,590,1448,182,PALE,round=True)
    d.text(105,613,1380,53,'最多100篇候选，共300段正文；输出最多3项',24,GREEN,True)
    d.text(105,684,1360,68,'模型只提供特征；自定义规则借鉴覆盖与MMR。不能把“换个说法”当成新知识。',20,INK)

    d.page('实际页面：保存中文目标并查看推荐',note='这是012实际应用截图，推荐来自生产worker与Spark。真实目标、推荐、阅读、主动已读与重算已走通。首轮存在偏题，重算语义partial，不能当质量通过。录制前预处理和等待压缩明示，不把缓存结果说成即时万条重算。')
    d.screenshot('01-workspace.png',72,223,1040,572)
    d.card(1142,232,382,544,'实际工程流程','中文目标真实保存\n\n生产Spark返回结果\n\n首轮偏题仍存在\n质量未通过')

    d.page('阅读与已读：行为改变可追踪',note='左侧为真实中文全文，右侧为主动已读记录和重算。已读数为1、修订号2的任务ready，已读Django事务材料被排除；语义状态partial。打开内容不会自动增加已读。工程闭环成功不代表学习效果已经证明。')
    d.screenshot('03-reading.png',72,232,748,560)
    d.screenshot('04-history.png',850,232,674,260)
    d.screenshot('05-recomputed.png',850,530,674,262)
    d.text(80,794,750,36,'原生中文全文阅读',14,GREEN,True)
    d.text(865,496,650,30,'显式已读记录',12,GREEN,True)
    d.text(865,796,650,30,'已读后的真实重算',12,GREEN,True)

    d.page('RSS：真实音频已经进入推荐，但闭环仍有限',note='使用AWS Morning Brief真实完整RSS条目和公开音频，ASR产生2961字符并实际入选。整个任务2013.8秒包括其他资料翻译，不是ASR耗时。播客中文未就绪，没有通知；冷启动入选不证明个人补充，也没有独立逐字金标。')
    d.card(76,237,695,474,'已经实际完成','真实RSS条目与公开音频\n完整音频进入ASR\n产生2961字符全文\n全文入库并实际入选')
    d.card(807,237,714,474,'尚未通过或未验证','播客中文本轮未就绪，零通知\n转写没有独立逐字金标准\n三分支补充通知仍待验证\n不能据首次失败说重试耗尽',WARN)
    d.text(84,747,1400,58,'RSS提供增量资料；最终仍须接受相同目标相关性和已读比较门槛。',20,INK)

    d.page('质量评价：保留真实失败，不把能运行当效果好',note='17个冻结场景中13个使用10215背景，4个受控。35个返回项逐条读原文，由代理审核，不是独立人类学习实验。13直接、13部分、9无关，即使宽松合并也低于80%。还发现11条错误补充、12条依据不足，核心质量未过。')
    d.text(80,229,1450,54,'17个冻结场景 · 35个返回项 · 四个技术主题',25,INK,True)
    segments=[(13,GREEN,'直接有帮助 13'),(13,'#92B9A0','部分相关 13'),(9,RED,'无关 9')]
    x=80
    for number,fill,label in segments:
        width=1440*number/35;d.rect(x,325,width,100,fill);d.text(x+15,351,width-30,57,label,22,WHITE,True);x+=width
    d.card(80,481,448,268,'37.1% 直接相关','宽松计部分相关为74.3%\n仍未达到冻结80%门槛',WARN)
    d.card(572,481,448,268,'11 条错误补充','另有12条解释依据不足\n同义复述误报仍存在',WARN)
    d.card(1064,481,448,268,'字面 ≠ 语义','348处原文位置校验通过\n不能抵消推荐含义错误')
    d.text(84,775,1440,40,'含明确受控场景；代理原文审核，不是学生学习提升实证。',15,MUTED)

    d.page('SDD与分工：需求、代码、证据能够对上',note='Spec-Kit先写规格与预期，再任务化并行开发。Git保留真实提交，不改写时间。提示词只导出可访问真实记录，注明范围。用户已确认最终分工：陶文杰负责数据算法，陈铭宇负责业务交互材料，各50%，与初版无调整；比例来自成员确认，不从AI代理工作量推算。')
    d.card(75,232,708,276,'可追溯过程','Spec / plan / tasks / contracts\n冻结输入 → 实现 → 原文核验\nGit真实提交 + 提示词来源归档')
    d.card(817,232,708,276,'最终分工（用户已确认）','陶文杰：数据、Spark、算法评价 50%\n陈铭宇：Mongo、API、Web/材料 50%\n与初版分工无调整')
    d.rect(75,550,1450,217,PALE,round=True)
    d.text(105,580,1370,52,'六项交付',24,GREEN,True)
    d.text(105,654,1370,102,'源码URL与Git · 需求/设计与提示词 · 项目报告\n≤5分钟MP4 · 分组分工表 · 可编辑汇报PPT',21,INK)

    d.page('下一步：先解决用户价值，再扩大规模',note='第一步提高正文和目标的直接相关性；第二步只在双方证据充分时说可能补充；第三步用用户要求的免费远程模型验证中文关键事实，并完成真实闭环。之后才做完整基线和扩大部署。当前12篇演示可展示系统，011仍保留未通过结论。')
    d.card(76,230,450,330,'01 找对材料','正文直接回答目标\n区分错误技术对象\n用独立原文事实验收')
    d.card(570,230,450,330,'02 解释对价值','同义复述不算新增\n双方证据不足就保留不确定\n完成正负通知对照')
    d.card(1064,230,450,330,'03 读懂中文','远程4次调用鉴权失败\n先解决HTTP 401，再核对语义\n失败不回退本地或收费')
    d.text(82,637,1440,105,'质量达到门槛以后，再决定增加语料、并发与跨主机Spark。',29,GREEN,True)
    d.text(85,769,1400,44,'当前结论：工程演示与质量验证分开，所有失败和依赖如实保留。',18,WARN,True)
    d.save()


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--draft',action='store_true')
    build(parser.parse_args().draft)
